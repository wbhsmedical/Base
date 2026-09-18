from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from engine.qa import check_pptx


def _pptx(slides: list[tuple[str, str]], tmp: Path) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for title, body in slides:
        s = prs.slides.add_slide(blank)
        t = s.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(12), Inches(1))
        t.text_frame.paragraphs[0].text = title
        b = s.shapes.add_textbox(Inches(0.5), Inches(1.6), Inches(12), Inches(4))
        b.text_frame.paragraphs[0].text = body
    p = tmp / "t.pptx"
    prs.save(p)
    return p


def test_ok(tmp_path):
    p = _pptx([("A", "one"), ("B", "two")], tmp_path)
    assert check_pptx(p) is None


def test_empty_slide(tmp_path):
    p = _pptx([("A", "one")], tmp_path)
    prs = Presentation(p)
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(p)
    err = check_pptx(p)
    assert err and "empty" in err


def test_lorem(tmp_path):
    p = _pptx([("A", "lorem ipsum dolor")], tmp_path)
    err = check_pptx(p)
    assert err and "placeholder" in err


def test_not_zip(tmp_path):
    p = tmp_path / "x.pptx"
    p.write_bytes(b"not a pptx")
    assert check_pptx(p)
