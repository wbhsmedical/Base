from __future__ import annotations

import pytest

from engine.job import Job
from engine.pptx_tools import NODE_MODULES, build
from engine.qa import check_pptx

PY_SCRIPT = r"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
s = prs.slides.add_slide(prs.slide_layouts[6])
b = s.shapes.add_textbox(Inches(0.5), Inches(1), Inches(12), Inches(2))
b.text_frame.paragraphs[0].text = "Hello deck"
b.text_frame.paragraphs[0].font.size = Pt(28)
prs.save(os.environ["PPTX_OUT"])
"""

JS_SCRIPT = r"""
const PptxGenJS = require("pptxgenjs");
const pptx = new PptxGenJS();
pptx.defineLayout({ name: "WIDE", width: 13.33, height: 7.5 });
pptx.layout = "WIDE";
const s = pptx.addSlide();
s.addText("Hello from js", { x: 0.6, y: 2, w: 12, h: 1.2, fontSize: 28 });
pptx.writeFile({ fileName: process.env.PPTX_OUT });
"""


@pytest.fixture
def job(tmp_path, monkeypatch):
    monkeypatch.setenv("BASE_DATA", str(tmp_path / "jobs"))
    return Job.create(model="local/fake", source_name="a.pdf", harness="loop")


def test_python_build(job):
    path, _ = build(job, PY_SCRIPT)
    assert check_pptx(path) is None


def test_js_build(job):
    if not (NODE_MODULES / "pptxgenjs").exists():
        pytest.skip("pptxgenjs not installed")
    job.backend = "js"
    path, _ = build(job, JS_SCRIPT)
    assert check_pptx(path) is None
